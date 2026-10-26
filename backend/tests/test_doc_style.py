"""Tests for the v2 doc-style schema and the DOCX renderer.

Run from ``backend/``:  python -m pytest tests -q
"""
import zipfile

import pytest
from docx import Document
from lxml import etree

from schemas.doc_style import STYLE_VERSION, StyleConfig
from services import docx_helpers as h
from services import docx_service

V1_CLASSIC = {
    "font_name": "Calibri", "font_size_name": 16, "font_size_section": 10,
    "font_size_body": 9.5, "font_size_contact": 9, "header_layout": "centered",
    "accent_color": "000000", "section_separator": "line", "name_bold": True,
    "name_uppercase": True, "section_caps": False, "margin_top": 0.5, "margin_bottom": 0.5,
    "margin_left": 0.5, "margin_right": 0.5, "space_before_section": 4, "space_after_section": 3,
    "section_heading_style": "line_below", "space_between_entries": 3.0,
    "entry_title_size": 10.0, "entry_subtitle_size": 9.5, "entry_subtitle_style": "bold",
    "entry_list_style": "bullet", "entry_indent_body": True, "color_dates": "555555",
    "experience_layout": "employer-title", "education_layout": "degree-school",
}

DATA = dict(
    user_name="Jane Doe", location="Austin, TX", email="jane@example.com", phone="555-0100",
    linkedin="linkedin.com/in/jane", summary="Engineer.",
    skills=[{"category": "Lang", "skills": ["Python", "Go"]}],
    educations=[{"school": "MIT", "degree": "BS", "field": "CS", "gpa": "3.9",
                 "start_date": "2008-09", "end_date": "2012-05"}],
    tailored_experiences=[{"company": "Acme", "location": "", "title": "Dev",
                           "start_date": "Jan 2018", "end_date": None, "bullets": ["Did a thing.", "Did more."]}],
    certifications=[{"name": "CKA", "issuer": "CNCF", "issue_date": "2021-03", "expiry_date": None}],
)


def test_v1_config_is_migrated():
    cfg = StyleConfig(**V1_CLASSIC)
    assert cfg.version == STYLE_VERSION
    assert cfg.fonts.family == "Calibri"
    assert cfg.page.margin_left == 0.5
    assert cfg.headings.para.borders.bottom.style == "single"
    assert cfg.entries.employer.bold and not cfg.entries.job_title.bold
    # certifications are appended to old section lists
    assert [s.key for s in cfg.sections][-1] == "certifications"


def test_legacy_separator_maps_to_borders():
    cfg = StyleConfig(**{**V1_CLASSIC, "section_separator": "double_line"})
    assert cfg.headings.para.borders.bottom.style == "double"
    cfg = StyleConfig(**{**V1_CLASSIC, "section_heading_style": "boxed"})
    b = cfg.headings.para.borders
    assert all(getattr(b, side).style == "single" for side in ("top", "left", "bottom", "right"))


def test_v2_roundtrip():
    cfg = StyleConfig(**V1_CLASSIC)
    again = StyleConfig(**cfg.model_dump())
    assert again == cfg


def test_invalid_colour_rejected():
    with pytest.raises(ValueError):
        StyleConfig(headings={"text": {"color": "not-a-colour"}})


@pytest.mark.parametrize("raw,fmt,expected", [
    ("2021-03", "mon_yyyy", "Mar 2021"),
    ("March 2021", "mm_yyyy", "03/2021"),
    ("03/2021", "month_yyyy", "March 2021"),
    ("Mar 2021", "yyyy", "2021"),
    ("garbage", "mon_yyyy", "garbage"),
    ("2021-03", "as_entered", "2021-03"),
    (None, "mon_yyyy", "Present"),
])
def test_format_date(raw, fmt, expected):
    assert h.format_date(raw, fmt) == expected


def _element_order_ok(xml_bytes: bytes) -> list[str]:
    """Return a list of schema-order violations in w:pPr / w:rPr children."""
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    root = etree.fromstring(xml_bytes)
    problems = []
    for tag, order in (("pPr", h._PPR_ORDER), ("rPr", h._RPR_ORDER)):
        names = [o.split(":")[1] for o in order]
        for el in root.iter(f"{{{ns['w']}}}{tag}"):
            idx = [names.index(etree.QName(c).localname) for c in el
                   if etree.QName(c).localname in names]
            if idx != sorted(idx):
                problems.append(f"{tag}: {[etree.QName(c).localname for c in el]}")
    return problems


FANCY = {
    "version": 2,
    "page": {"size": "a4", "page_border": {"style": "single", "width": 1, "color": "2E74B5"}},
    "fonts": {"family": "Georgia", "kerning": True, "auto_hyphenation": True},
    "header": {"layout": "left", "name": {"size": 20, "small_caps": True, "letter_spacing": 1.5,
                                           "highlight": "FFFF00", "underline": "wavy"},
               "link_underline": True,
               "contact_para": {"borders": {"bottom": {"style": "double", "width": 1.5}}}},
    "headings": {"text": {"color": "FFFFFF", "all_caps": True},
                 "para": {"shading": "1F3864", "borders": {"left": {"style": "single", "width": 3}}}},
    "entries": {"experience_layout": "combined", "date_format": "mon_yyyy",
                "bullets": {"kind": "numbered"}},
    "header_footer": {"header_text": "{name}", "page_numbers": "page_of_total",
                      "different_first_page": True},
    "cover_letter": {"letterhead": "match_resume", "include_recipient": True,
                     "header_footer": {"page_numbers": "number"}},
}


@pytest.fixture(scope="module")
def rendered(tmp_path_factory):
    d = tmp_path_factory.mktemp("docs")
    style = StyleConfig(**FANCY)
    resume = docx_service.create_resume(**DATA, output_path=d / "r.docx", style=style)
    letter = docx_service.create_cover_letter(
        user_name="Jane Doe", email="jane@example.com", phone="555-0100",
        cover_letter_text="Dear X,\n\nHello.\n\nBest,\nJane", job_title="Dev", company="Globex",
        output_path=d / "c.docx", style=style,
    )
    return resume, letter


def test_default_style_renders(tmp_path):
    out = docx_service.create_resume(**DATA, output_path=tmp_path / "r.docx")
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "JANE DOE" in text and "Certifications" in text and "CKA" in text


def test_legacy_style_renders_same_text(tmp_path):
    out = docx_service.create_resume(**DATA, output_path=tmp_path / "r.docx", style=StyleConfig(**V1_CLASSIC))
    text = "\n".join(p.text for p in Document(str(out)).paragraphs)
    assert "Professional Experience" in text and "Jan 2018 – Present" in text


def test_xml_element_order(rendered):
    for path in rendered:
        with zipfile.ZipFile(path) as z:
            for part in z.namelist():
                if part.startswith("word/") and any(k in part for k in ("document", "header", "footer")):
                    assert _element_order_ok(z.read(part)) == [], part


def test_features_present_in_xml(rendered):
    resume, letter = rendered
    with zipfile.ZipFile(resume) as z:
        doc = z.read("word/document.xml").decode()
        names = z.namelist()
        assert "w:pgBorders" in doc and "w:hyperlink" in doc and "w:numPr" in doc
        assert 'w:pgSz w:w="11905" w:h="16838"' in doc  # A4 in twips
        assert "w:smallCaps" in doc and "w:shd" in doc and "w:kern" in doc
        assert any(n.startswith("word/header") for n in names)
        assert any(n.startswith("word/footer") for n in names)
        assert "<w:titlePg" in doc
        assert "w:autoHyphenation" in z.read("word/settings.xml").decode()
        assert "decimal" in z.read("word/numbering.xml").decode()
    with zipfile.ZipFile(letter) as z:
        assert "Hiring Manager" in z.read("word/document.xml").decode()


def test_sections_hidden_and_reordered(tmp_path):
    style = StyleConfig(sections=[{"key": "education"}, {"key": "summary", "visible": False}])
    out = docx_service.create_resume(**DATA, output_path=tmp_path / "r.docx", style=style)
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert "Professional Summary" not in texts
    assert texts.index("Education") < texts.index("Professional Experience")
